#!/usr/bin/env python3
"""Build the detailed product specification DOCX from the checked-in Markdown."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "郑育星本地教育服务小程序完整产品文档_V1.0_2026-07-28.md"
OUTPUT = ROOT / "docs" / "郑育星本地教育服务小程序完整产品文档_V1.0_2026-07-28.docx"

NAVY = "1E3A5F"
TEAL = "2BB3A3"
TEXT = "1F2933"
MUTED = "6B7280"
BORDER = "D8E0E8"
HEADER_FILL = "E8EEF5"
SOFT_FILL = "F6F8FA"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = BORDER, size: str = "4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, name="Arial Unicode MS", east_asia="Arial Unicode MS", size=None, color=None, bold=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def add_inline_runs(paragraph, text: str, size=None, color=TEXT, bold=False, code=False):
    token_re = re.compile(r"(\*\*.+?\*\*|`.+?`)")
    position = 0
    for match in token_re.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            set_run_font(run, size=size, color=color, bold=bold)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, color=color, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Arial Unicode MS", east_asia="Arial Unicode MS", size=size or 9, color=NAVY)
            run.font.highlight_color = None
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size=size, color=color, bold=bold)


def set_paragraph_borders(paragraph, fill=SOFT_FILL, color=BORDER):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "4")
        node.set(qn("w:color"), color)
        p_bdr.append(node)
    p_pr.append(p_bdr)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial Unicode MS"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_specs = {
        "Heading 1": (16, NAVY, 16, 6),
        "Heading 2": (13, NAVY, 13, 4),
        "Heading 3": (12, "284766", 10, 3),
        "Heading 4": (11, TEXT, 8, 2),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = styles[name]
        style.font.name = "Arial Unicode MS"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Number"):
        style = styles[list_name]
        style.font.name = "Arial Unicode MS"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.2


def add_field(paragraph, field_code: str) -> None:
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = field_code
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_end])


def add_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("郑育星  |  产品需求与实现基线")
    set_run_font(run, size=8.5, color=MUTED)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("内部交付文档  ·  V1.0  ·  ")
    set_run_font(run, size=8.5, color=MUTED)
    add_field(p, "PAGE")


def add_cover(doc: Document, metadata: list[list[str]]) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(65)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run("郑育星")
    set_run_font(run, size=30, color=TEAL, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(24)
    run = p.add_run("本地教育服务小程序完整产品文档")
    set_run_font(run, size=22, color=NAVY, bold=True)

    accent = doc.add_paragraph()
    accent.paragraph_format.space_after = Pt(28)
    run = accent.add_run("产品需求  ·  当前实现  ·  上线验收")
    set_run_font(run, size=11, color=MUTED)
    p_pr = accent._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "16")
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), TEAL)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for left, right in metadata:
        cells = table.add_row().cells
        cells[0].width = Inches(1.45)
        cells[1].width = Inches(5.05)
        cells[0].text = left
        cells[1].text = right
        set_cell_shading(cells[0], HEADER_FILL)
        set_cell_shading(cells[1], WHITE)
        for index, cell in enumerate(cells):
            set_cell_border(cell)
            set_cell_margins(cell)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                for run in para.runs:
                    set_run_font(run, size=9.5, color=NAVY if index == 0 else TEXT, bold=index == 0)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(22)
    set_paragraph_borders(p, fill="EFF8F6", color="BCE7DF")
    add_inline_runs(
        p,
        "当前结论：项目已具备本地可运行 MVP，但真实微信支付、生产数据库、正式域名与密钥、上传存储、后台完整管理和真机验收仍是上线前置条件。",
        size=10,
        color=TEXT,
        bold=False,
    )


def add_contents(doc: Document, chapter_titles: list[str]) -> None:
    heading = doc.add_paragraph(style="Heading 1")
    heading.paragraph_format.page_break_before = True
    heading.add_run("目录")
    intro = doc.add_paragraph()
    add_inline_runs(intro, "本目录按产品、业务、工程与上线验收顺序组织。", size=9.5, color=MUTED)
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    midpoint = (len(chapter_titles) + 1) // 2
    left = chapter_titles[:midpoint]
    right = chapter_titles[midpoint:]
    for idx in range(midpoint):
        cells = table.add_row().cells
        values = [left[idx], right[idx] if idx < len(right) else ""]
        for cell, value in zip(cells, values):
            cell.width = Inches(3.25)
            set_cell_border(cell, color=WHITE, size="0")
            set_cell_margins(cell, top=50, bottom=50)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline_runs(p, value, size=9.5, color=NAVY)
    doc.add_page_break()


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=0, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    font_size = 8.2 if cols >= 5 else 8.8 if cols == 4 else 9.3
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        if row_index == 0:
            set_repeat_table_header(table.rows[-1])
        for col_index, cell in enumerate(cells):
            value = values[col_index] if col_index < len(values) else ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_border(cell)
            set_cell_shading(cell, HEADER_FILL if row_index == 0 else WHITE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            add_inline_runs(
                p,
                value,
                size=font_size,
                color=NAVY if row_index == 0 else TEXT,
                bold=row_index == 0,
            )
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        values = [part.strip() for part in lines[index].strip().strip("|").split("|")]
        if index == start + 1 and all(re.fullmatch(r":?-{3,}:?", value) for value in values):
            index += 1
            continue
        rows.append(values)
        index += 1
    return rows, index


def build() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)
    add_header_footer(doc)

    metadata = []
    first_table_start = next(i for i, line in enumerate(lines) if line.strip().startswith("|"))
    metadata_rows, after_metadata = parse_table(lines, first_table_start)
    for row in metadata_rows[1:]:
        if len(row) >= 2:
            metadata.append([row[0], row[1]])

    add_cover(doc, metadata)
    chapter_titles = [line[3:].strip() for line in lines if line.startswith("## ")]
    add_contents(doc, chapter_titles)

    index = after_metadata
    in_code = False
    code_lines: list[str] = []
    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()

        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.18)
                p.paragraph_format.right_indent = Inches(0.18)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(7)
                set_paragraph_borders(p, fill="F3F5F7", color="D7DEE5")
                run = p.add_run("\n".join(code_lines))
                set_run_font(run, name="Arial Unicode MS", east_asia="Arial Unicode MS", size=8.5, color=TEXT)
                in_code = False
            index += 1
            continue

        if in_code:
            code_lines.append(raw)
            index += 1
            continue

        if not stripped:
            index += 1
            continue

        if stripped.startswith("|"):
            table_rows, index = parse_table(lines, index)
            add_markdown_table(doc, table_rows)
            continue

        if raw.startswith("# "):
            index += 1
            continue
        if raw.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_inline_runs(p, raw[3:].strip(), size=16, color=NAVY, bold=True)
            index += 1
            continue
        if raw.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline_runs(p, raw[4:].strip(), size=13, color=NAVY, bold=True)
            index += 1
            continue
        if raw.startswith("#### "):
            p = doc.add_paragraph(style="Heading 3")
            add_inline_runs(p, raw[5:].strip(), size=12, color="284766", bold=True)
            index += 1
            continue
        if raw.startswith("##### "):
            p = doc.add_paragraph(style="Heading 4")
            add_inline_runs(p, raw[6:].strip(), size=11, color=TEXT, bold=True)
            index += 1
            continue

        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.right_indent = Inches(0.12)
            set_paragraph_borders(p, fill="EFF8F6", color="BCE7DF")
            add_inline_runs(p, stripped.lstrip("> ").strip(), size=10, color=TEXT)
            index += 1
            continue

        if re.match(r"^[-*]\s+", stripped):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, re.sub(r"^[-*]\s+", "", stripped), size=10.5, color=TEXT)
            index += 1
            continue

        numbered_match = re.match(r"^(\d+)\.\s+(.*)", stripped)
        if numbered_match:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.375)
            p.paragraph_format.first_line_indent = Inches(-0.188)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.2
            prefix = p.add_run(f"{numbered_match.group(1)}. ")
            set_run_font(prefix, size=10.5, color=TEXT)
            add_inline_runs(p, numbered_match.group(2), size=10.5, color=TEXT)
            index += 1
            continue

        p = doc.add_paragraph()
        add_inline_runs(p, stripped, size=10.5, color=TEXT)
        index += 1

    core_props = doc.core_properties
    core_props.title = "郑育星本地教育服务小程序完整产品文档"
    core_props.subject = "产品需求、当前实现、接口、数据与上线验收"
    core_props.author = "郑育星项目组"
    core_props.keywords = "微信小程序, 家长, 大学生老师, 预约, 解锁, 订单, 后台"
    core_props.comments = "基于 2026-07-28 当前代码和需求基线生成"

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
