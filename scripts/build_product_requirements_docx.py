from pathlib import Path
import re
import argparse

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "PRODUCT_REQUIREMENTS.md"
OUTPUT = ROOT / "docs" / "郑州大学生家教小程序产品文档.docx"
DEFAULT_TITLE = "郑州大学生家教小程序"
DEFAULT_SUBTITLE = "MVP 产品定位、需求范围、业务规则与数据埋点方案"

NAVY = "1E3A5F"
TEAL = "2BB3A3"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2933"
MUTED = "6B7280"
LIGHT_BLUE = "E8EEF5"
LIGHT_TEAL = "E8F7F4"
LIGHT_GRAY = "F2F4F7"
BORDER = "D7E0EA"
WHITE = "FFFFFF"

CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_run_font(run, size=None, color=None, bold=None, italic=None, font_name="Calibri"):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:ascii"), font_name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font_name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for grid_col in list(grid):
        grid.remove(grid_col)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color=BORDER):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_para_border(paragraph, color=TEAL, width="18", space="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), width)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_text(paragraph, text, size=11, color=INK, bold=False, italic=False):
    """Render simple Markdown bold/code tokens while keeping a clean Word paragraph."""
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, size=size, color=color, bold=bold, italic=italic)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, color=color, bold=True, italic=italic)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=max(size - 0.5, 9), color=DARK_BLUE, bold=True, font_name="Consolas")
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=size, color=color, bold=bold, italic=italic)


def add_body_paragraph(doc, text, style="Normal"):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    add_text(p, text)
    return p


def add_list_item(doc, text, numbered=False):
    p = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    add_text(p, text)
    return p


def set_keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def add_heading(doc, level, text):
    style = f"Heading {level}"
    p = doc.add_paragraph(style=style)
    if level == 1:
        p.paragraph_format.page_break_before = False
    set_keep_with_next(p)
    add_text(p, text, size={1: 16, 2: 13, 3: 12}[level], color={1: BLUE, 2: BLUE, 3: DARK_BLUE}[level], bold=True)
    return p


def split_table_row(line):
    return [part.strip().replace("\\|", "|") for part in line.strip().strip("|").split("|")]


def is_separator_row(line):
    stripped = line.replace("|", "").replace("-", "").replace(":", "").replace(" ", "")
    return not stripped


def add_markdown_table(doc, rows):
    data = [split_table_row(row) for row in rows if not is_separator_row(row)]
    if not data:
        return
    cols = max(len(row) for row in data)
    data = [row + [""] * (cols - len(row)) for row in data]
    table = doc.add_table(rows=len(data), cols=cols)
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_table_borders(table)

    if cols == 2:
        widths = [2700, 6660]
    elif cols == 3:
        widths = [2200, 3600, 3560]
    elif cols == 4:
        widths = [2200, 2700, 2700, 1760]
    else:
        base = CONTENT_WIDTH_DXA // cols
        widths = [base] * cols
        widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    set_table_geometry(table, widths)

    for row_index, values in enumerate(data):
        row = table.rows[row_index]
        if row_index == 0:
            set_repeat_table_header(row)
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            if row_index == 0:
                shade_cell(cell, LIGHT_BLUE)
            elif row_index % 2 == 0:
                shade_cell(cell, "FAFBFC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            add_text(p, value, size=9.5, color=INK, bold=(row_index == 0))
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_code_box(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    set_table_borders(table, color="C8D9D5")
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    cell = table.cell(0, 0)
    shade_cell(cell, LIGHT_TEAL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.1
    run = p.add_run("\n".join(lines))
    set_run_font(run, size=9.5, color=DARK_BLUE, font_name="Consolas")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for level, size, color, before, after in [
        (1, 16, BLUE, 18, 10),
        (2, 13, BLUE, 14, 7),
        (3, 12, DARK_BLUE, 10, 5),
    ]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0

    for style_name in ["List Bullet", "List Number"]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    cover_title = doc.styles.add_style("Cover Title", WD_STYLE_TYPE.PARAGRAPH)
    cover_title.font.name = "Calibri"
    cover_title._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    cover_title.font.size = Pt(27)
    cover_title.font.bold = True
    cover_title.font.color.rgb = RGBColor.from_string(NAVY)
    cover_title.paragraph_format.space_before = Pt(0)
    cover_title.paragraph_format.space_after = Pt(8)

    cover_subtitle = doc.styles.add_style("Cover Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    cover_subtitle.font.name = "Calibri"
    cover_subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    cover_subtitle.font.size = Pt(13)
    cover_subtitle.font.color.rgb = RGBColor.from_string(MUTED)
    cover_subtitle.paragraph_format.space_after = Pt(20)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    suffix = paragraph.add_run(" 页")
    set_run_font(suffix, size=9, color=MUTED)


def add_header_footer(section, title):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(f"{title}  |  产品文档")
    set_run_font(run, size=9, color=MUTED, bold=True)
    set_para_border(p, color="D9E4EE", width="6", space="5")

    footer = section.footer
    p = footer.paragraphs[0]
    add_page_number(p)


def add_cover(doc, title, subtitle):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("产品需求文档")
    set_run_font(run, size=11, color=TEAL, bold=True)

    p = doc.add_paragraph(style="Cover Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_text(p, title, size=27, color=NAVY, bold=True)

    p = doc.add_paragraph(style="Cover Subtitle")
    add_text(p, subtitle, size=13, color=MUTED)
    set_para_border(p, color=TEAL, width="20", space="12")

    table = doc.add_table(rows=5, cols=2)
    table.autofit = False
    set_table_borders(table)
    set_table_geometry(table, [2500, 6860])
    rows = [
        ("文档范围", subtitle),
        ("服务城市", "郑州市"),
        ("产品阶段", "MVP（最小可行产品）"),
        ("目标用户", "家长、郑州高校大学生老师、平台管理员"),
        ("文档版本", "V1.0"),
    ]
    for index, values in enumerate(rows):
        for col, value in enumerate(values):
            cell = table.cell(index, col)
            shade_cell(cell, LIGHT_TEAL if col == 0 else WHITE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_text(p, value, size=10.5, color=NAVY if col == 0 else INK, bold=(col == 0))

    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    note_table = doc.add_table(rows=1, cols=1)
    note_table.autofit = False
    set_table_borders(note_table, color="C9DED9")
    set_table_geometry(note_table, [CONTENT_WIDTH_DXA])
    note_cell = note_table.cell(0, 0)
    shade_cell(note_cell, "F5FBFA")
    p = note_cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_text(p, "文档说明：本文件用于产品、研发、测试和运营团队对齐 MVP 范围、业务流程、权限边界与数据分析口径。", size=10.5, color=DARK_BLUE)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    toc_heading = doc.add_paragraph()
    toc_heading.paragraph_format.space_before = Pt(18)
    toc_heading.paragraph_format.space_after = Pt(8)
    add_text(toc_heading, "文档结构", size=13, color=BLUE, bold=True)
    for item in ["项目背景", "需求列举", "需求描述", "建议埋点"]:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        add_text(p, item, size=10.5, color=INK)

    doc.add_page_break()


def parse_markdown_into_doc(doc, text):
    lines = text.splitlines()
    # Exclude title and metadata table because the cover has a formatted equivalent.
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        if index <= 9:
            index += 1
            continue
        if stripped.startswith("```"):
            code_lines = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            add_code_box(doc, code_lines)
            index += 1
            continue
        if stripped.startswith("|"):
            table_rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_rows.append(lines[index])
                index += 1
            add_markdown_table(doc, table_rows)
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            text_value = heading.group(2)
            if level == 1:
                # Markdown root title is intentionally replaced by cover.
                index += 1
                continue
            add_heading(doc, level - 1, text_value)
            index += 1
            continue
        bullet = re.match(r"^-\s+(.+)$", stripped)
        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if bullet:
            add_list_item(doc, bullet.group(1), numbered=False)
            index += 1
            continue
        if numbered:
            add_list_item(doc, numbered.group(1), numbered=True)
            index += 1
            continue
        add_body_paragraph(doc, stripped)
        index += 1


def build(source=SOURCE, output=OUTPUT, title=DEFAULT_TITLE, subtitle=DEFAULT_SUBTITLE):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(doc)
    add_header_footer(section, title)
    add_cover(doc, title, subtitle)
    parse_markdown_into_doc(doc, source.read_text(encoding="utf-8"))
    doc.core_properties.title = title
    doc.core_properties.subject = subtitle
    doc.core_properties.author = "郑州大学生家教项目组"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", default=str(SOURCE))
    parser.add_argument("--output", default=str(OUTPUT))
    parser.add_argument("--title", default=DEFAULT_TITLE)
    parser.add_argument("--subtitle", default=DEFAULT_SUBTITLE)
    args = parser.parse_args()
    build(Path(args.source), Path(args.output), args.title, args.subtitle)
